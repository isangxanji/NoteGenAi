
import csv
from io import StringIO
from pathlib import Path

from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook


MAX_EXTRACTED_CHARACTERS = 500_000


def extract_text(file_path, file_type):
    file_type = file_type.lower().lstrip(".")

    if file_type == "pdf":
        reader = PdfReader(file_path)

        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are not supported.")

        text_parts = []

        for page in reader.pages:
            text_parts.append(page.extract_text() or "")

        extracted_text = "\n".join(text_parts)

    elif file_type == "txt":
        try:
            with open(file_path, "r", encoding="utf-8-sig") as file:
                extracted_text = file.read()
        except UnicodeDecodeError:
            with open(file_path, "r", encoding="cp1252") as file:
                extracted_text = file.read()

    elif file_type == "docx":
        document = Document(file_path)
        parts = [paragraph.text for paragraph in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        extracted_text = "\n".join(parts)

    elif file_type == "pptx":
        presentation = Presentation(file_path)
        parts = []

        for slide_number, slide in enumerate(presentation.slides, start=1):
            parts.append(f"Slide {slide_number}")

            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text)

                if shape.has_table:
                    for row in shape.table.rows:
                        parts.append(
                            " | ".join(cell.text for cell in row.cells)
                        )

        extracted_text = "\n".join(parts)

    elif file_type == "xlsx":
        workbook = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True
        )
        parts = []

        try:
            for sheet in workbook.worksheets:
                parts.append(f"Sheet: {sheet.title}")

                for row in sheet.iter_rows(values_only=True):
                    values = [
                        str(value) if value is not None else ""
                        for value in row
                    ]

                    if any(values):
                        parts.append(" | ".join(values))
        finally:
            workbook.close()

        extracted_text = "\n".join(parts)

    elif file_type == "csv":
        try:
            with open(
                file_path, "r", encoding="utf-8-sig", newline=""
            ) as file:
                reader = csv.reader(file)
                extracted_text = "\n".join(
                    " | ".join(row) for row in reader
                )
        except UnicodeDecodeError:
            with open(
                file_path, "r", encoding="cp1252", newline=""
            ) as file:
                reader = csv.reader(file)
                extracted_text = "\n".join(
                    " | ".join(row) for row in reader
                )

    else:
        raise ValueError("This file type is not supported.")

    extracted_text = extracted_text.strip()

    if not extracted_text:
        raise ValueError(
            "No readable text was found. The file may be empty "
            "or may need OCR."
        )

    if len(extracted_text) > MAX_EXTRACTED_CHARACTERS:
        raise ValueError(
            "This document contains too much text for this stage. "
            "Please try a smaller document."
        )

    return extracted_text
